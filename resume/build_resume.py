"""Builds Evan_Borden_Resume.docx (single-column, ATS-friendly).

Run with a Python that has python-docx installed:
    python3 -m venv .venv && .venv/bin/pip install python-docx
    .venv/bin/python resume/build_resume.py
Then export the PDF from Word (File > Save As > PDF) or via export_pdf.sh.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name("Evan_Borden_Resume.docx")

FONT = "Calibri"
INK = RGBColor(0x1F, 0x23, 0x28)
ACCENT = RGBColor(0x1F, 0x4E, 0x5F)
MUTED = RGBColor(0x55, 0x5B, 0x63)
MARGIN = Inches(0.7)
TEXT_WIDTH = Inches(8.5) - 2 * MARGIN

# ---------------------------------------------------------------- content

NAME = "Evan Borden"
HEADLINE = "Manager of Engineering  ·  Adobe Architect"
# (label, link or None)
CONTACT = [
    ("Charlotte, NC", None),
    ("704.401.4864", "tel:+17044014864"),
    ("evanpatrickborden@outlook.com", "mailto:evanpatrickborden@outlook.com"),
    ("linkedin.com/in/evan-borden", "https://www.linkedin.com/in/evan-borden/"),
    ("evanborden.com", "https://www.evanborden.com/"),
]

SUMMARY = (
    "Engineering manager and Adobe architect at Razorfish, a martech agency. I lead the people and the plan "
    "behind client teams (hiring, staffing and resourcing development and QA engineers across the US, India "
    "and Costa Rica) and own the architecture those teams build, most recently Adobe Experience Platform and "
    "Journey Optimizer work in healthcare under HIPAA. Trusted to run more than one flagship account at once, "
    "on a hands-on engineering foundation across Adobe Experience Manager, the Adobe cloud stack and "
    "enterprise WordPress."
)

SKILLS = [
    ("Leadership & Delivery",
     "Technical hiring and interviewing · Global talent sourcing (India, Costa Rica) · Resource planning and margin · "
     "SOWs, RFPs and estimation · Architecture review boards · Client communication and live demos"),
    ("Adobe Experience Cloud",
     "AEM as a Cloud Service · AEM 6.5 · App Builder · Adobe I/O · Experience Platform · Journey Optimizer · "
     "Offer Decisioning · XDM · Data Collection · Target · Analytics and CJA"),
    ("Engineering",
     "Java · JavaScript · React · Vue · Node.js · PHP · WordPress and Gutenberg · REST API design · MySQL · "
     "CI/CD (GitLab, Azure DevOps) · Docker · Linux · AWS · Azure"),
    ("Security & AI",
     "HIPAA and PHI data-flow review · OAuth/OIDC · JWT/JOSE · WCAG · Claude and ChatGPT · Jira and Confluence automation"),
]

RAZORFISH_BULLETS = [
    "Own hiring, staffing and interviewing for development and QA roles across experience levels and role types.",
    "Partner with staffing leads to source engineers in Costa Rica and India, filling project roles from the internal "
    "bench and external candidates.",
    "Work with project managers on resource start and end dates, role rates and cost, adjusting staffing to protect margin.",
    "Contribute directly to statements of work, review RFPs, and vet technical requirements and levels of effort with "
    "delivery teams.",
    "Take architecture recommendations through client architecture review boards, tracing every data path for "
    "security and PHI exposure under HIPAA.",
    "Build slide decks, demo technical functionality to live audiences, and write architecture research and "
    "discovery documentation.",
    "Use Claude, ChatGPT and other models to draft architecture documents, write Confluence documentation against real "
    "codebases, automate Jira ticket resolution, and generate Word, Excel and PowerPoint deliverables.",
    "Advocate for engineers' promotions and surface their work to leadership; standardized teams on shared coding "
    "conventions.",
]

ACCOUNTS = [
    ("Labcorp", "Adobe Architect & Tech Lead", "2026 – Present", [
        "Entrusted with leading the account by the GVP of Technology and CTO at the time. Took both 2026 projects into "
        "production, which led to continued work with Labcorp.",
        "Thrive 5 Personalization: owned the solution architecture across AEMaaCS, App Builder, AEP/AJO and a downstream "
        "HIPAA-compliant Patient Portal. My team built the App Builder integration tier and the AEP/AJO layer (XDM "
        "schemas, profiles, journeys, offer decisioning); I coordinated partner teams and client stakeholders through "
        "architecture review board sign-off.",
        "Marker by Labcorp: led the technical implementation and the team building custom AEMaaCS components, back end "
        "and front end, for a media-driven product landing page on a tight timeline, on camera with the client throughout.",
    ]),
    ("Disney Rewards", "Tech Lead", "2022 – 2026", [
        "Tech Lead on a custom WordPress platform after developing on it since 2019: bespoke Gutenberg blocks, ACF Pro "
        "architecture, Chase Bank API integration via JWT/JOSE, Adobe CJA and Adobe Target, shipped through GitLab CI/CD. "
        "Hired my own successor when I moved to lead Labcorp.",
    ]),
    ("UC Health", "Tech Lead & Primary AEM Developer", "2018 – 2023", [
        "Led and built an on-premise AEM 6.5 platform for a large healthcare system, in parallel with Disney Rewards: "
        "multi-module Maven build, OSGi services, Sling models and servlets, HTL, Vue 2 clientlibs, and Epic EHR "
        "integration architecture.",
    ]),
]

INTERACTIVE_KNOWLEDGE = [
    "Built media-rich interactive experiences for non-profits, museums and educational institutions.",
    "Shipped responsive, mobile-first apps in React and Vue against custom REST and microservice backends.",
    "Partnered with UX to translate design comps into accessible, performant interfaces.",
]

EDUCATION = [
    ("UNC Charlotte", "B.S. Computer Science, Software Engineering Concentration"),
    ("Central Piedmont Community College", "A.A.S. Advertising and Graphic Design"),
]

# ---------------------------------------------------------------- helpers

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = MARGIN
sec.top_margin = sec.bottom_margin = Inches(0.6)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.08

doc.core_properties.author = NAME
doc.core_properties.title = f"{NAME} — Resume"


def para(space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def run(p, text, size=None, bold=False, italic=False, color=None, caps=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if caps:
        r.font.all_caps = True
    return r


def link(p, text, url, size=None, color=None):
    """Clickable hyperlink run (kept in the PDF when Word exports it)."""
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), str(color))
        rPr.append(c)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def bottom_rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "2")
    b.set(qn("w:color"), "1F4E5F")
    bdr.append(b)
    pPr.append(bdr)


def heading(text):
    p = para(space_before=11, space_after=5)
    run(p, text, size=11, bold=True, color=ACCENT, caps=True)
    bottom_rule(p)
    p.paragraph_format.keep_with_next = True


def role_line(left_bold, left_rest, right, space_before=6, size=11, indent=0):
    p = para(space_before=space_before, space_after=1)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    run(p, left_bold, size=size, bold=True)
    if left_rest:
        run(p, f"  |  {left_rest}", size=size - 0.5, color=MUTED)
    run(p, f"\t{right}", size=10, color=MUTED)
    return p


def title_line(title, dates):
    """Stacked job title under a company line (shows promotions at one employer)."""
    p = para(space_after=1)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    run(p, title, size=10.5, color=INK)
    run(p, f"\t{dates}", size=10, color=MUTED)
    return p


def bullet(text, indent=0.18):
    """Bullet with a real gap after the dot. A leading 'Label: ' is bolded."""
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    hang = Inches(0.16)
    pf.left_indent = Inches(indent) + hang
    pf.first_line_indent = -hang
    pf.tab_stops.add_tab_stop(Inches(0.25), WD_TAB_ALIGNMENT.CLEAR)  # List Bullet's built-in tab
    pf.tab_stops.add_tab_stop(Inches(indent) + hang)
    pf.space_after = Pt(2)
    label, sep, rest = text.partition(": ")
    if sep and len(label) < 32:
        run(p, label + ": ", bold=True)
        run(p, rest)
    else:
        run(p, text)
    return p


# ---------------------------------------------------------------- layout

p = para()
run(p, NAME, size=24, bold=True, color=INK)
p = para(space_after=3)
run(p, HEADLINE, size=12, color=ACCENT)
p = para(space_after=2)
for i, (label, url) in enumerate(CONTACT):
    if i:
        run(p, "  ·  ", size=9.5, color=MUTED)
    if url:
        link(p, label, url, size=9.5, color=MUTED)
    else:
        run(p, label, size=9.5, color=MUTED)

heading("Summary")
p = para()
run(p, SUMMARY)

heading("Core Skills")
for label, items in SKILLS:
    p = para(space_after=2)
    run(p, f"{label}:  ", bold=True)
    run(p, items)

heading("Experience")
role_line("Razorfish", None, "Nov 2018 – Present", space_before=2)
title_line("Manager of Engineering", "2024 – Present")
title_line("Senior Engineer", "Nov 2018 – 2024")
for b in RAZORFISH_BULLETS:
    bullet(b)

p = para(space_before=5, space_after=1)
p.paragraph_format.left_indent = Inches(0.18)
run(p, "Key accounts", size=10, bold=True, italic=True, color=MUTED)
for account, title, dates, items in ACCOUNTS:
    role_line(account, title, dates, space_before=4, size=10.5, indent=0.18)
    for b in items:
        bullet(b, indent=0.36)

role_line("Interactive Knowledge", "Senior Web Developer", "Sep 2012 – Oct 2018", space_before=9)
for b in INTERACTIVE_KNOWLEDGE:
    bullet(b)

role_line("Earlier Roles", "Web Developer", "2008 – 2012", space_before=9)
bullet("Web development roles that built the technical foundation for everything above.")

heading("Education")
for school, degree in EDUCATION:
    p = para(space_after=2)
    run(p, school, bold=True)
    run(p, f"  |  {degree}", color=MUTED)

doc.save(OUT)
print(f"wrote {OUT}")
